import os
import json
from fastapi import APIRouter, Request, Form, UploadFile, File
from fastapi.responses import RedirectResponse
from fastapi.templating import Jinja2Templates
from app.auth import login_required, can_edit, notify_other_editors
from app.database import get_db
from app.config import PER_PAGE, PHOTO_DIR

router = APIRouter(prefix="/persons")
templates = Jinja2Templates(directory="app/templates")

INDUSTRY_ORDER = ['品牌', '餐飲', '銷售']

def industry_priority(tag):
    if not tag:
        return 9999
    for i, kw in enumerate(INDUSTRY_ORDER):
        if kw in tag:
            return i
    return len(INDUSTRY_ORDER)

TITLE_RANKS = [
    (1, ['董事長','主席','創辦人','共同創辦人','執行長','總裁','founder','chairman','ceo','president']),
    (2, ['coo','cfo','cto','cmo']),
    (3, ['總經理','副總裁','evp','svp']),
    (4, ['副總','執行副總','vp','副總經理']),
    (5, ['協理','總監','director']),
    (6, ['經理','manager']),
    (7, ['副理','主任','supervisor']),
    (8, ['組長','課長','leader']),
    (9, ['專員','工程師','設計師','analyst','associate']),
]

def title_priority(title):
    if not title:
        return 999
    t = title.lower()
    for level, keys in TITLE_RANKS:
        if any(k in t for k in keys):
            return level
    return 50


def _render(request, template, **kwargs):
    ctx = {"request": request, "user": request.state.user,
           "get_flashed_messages": lambda: request.state.flash,
           "notif_count": getattr(request.state, "notif_count", 0),
           "can_edit": can_edit(request.state.user), **kwargs}
    return templates.TemplateResponse(request=request, name=template, context=ctx)


def _get_or_create_tag(db, name: str) -> int:
    name = name.strip().lower()
    row = db.execute("SELECT id FROM tags WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    db.execute("INSERT INTO tags (name) VALUES (?)", (name,))
    db.commit()
    return db.execute("SELECT last_insert_rowid()").fetchone()[0]


def _save_tags(db, person_id: int, tags_str: str):
    db.execute("DELETE FROM person_tags WHERE person_id = ?", (person_id,))
    if tags_str:
        for tag_name in tags_str.split(","):
            tag_name = tag_name.strip()
            if tag_name:
                tag_id = _get_or_create_tag(db, tag_name)
                db.execute("INSERT OR IGNORE INTO person_tags (person_id, tag_id) VALUES (?, ?)",
                           (person_id, tag_id))
    db.commit()


def _save_roles(db, person_id: int, form_data):
    db.execute("DELETE FROM roles WHERE person_id = ?", (person_id,))

    titles = form_data.getlist("role_title[]")
    company_names = form_data.getlist("role_company_name[]")
    work_emails = form_data.getlist("role_work_email[]")
    work_phones = form_data.getlist("role_work_phone[]")
    current_checks = form_data.getlist("role_is_current[]")

    for i in range(len(titles)):
        company_name = company_names[i].strip() if i < len(company_names) else ""
        company_id = None
        if company_name:
            existing = db.execute("SELECT id FROM companies WHERE name = ?", (company_name,)).fetchone()
            if existing:
                company_id = existing["id"]
            else:
                db.execute("INSERT INTO companies (name) VALUES (?)", (company_name,))
                db.commit()
                company_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]

        title = titles[i] if i < len(titles) else ""
        work_email = work_emails[i] if i < len(work_emails) else ""
        work_phone = work_phones[i] if i < len(work_phones) else ""
        is_current = str(i) in current_checks or str(i + 1) in current_checks

        if title or company_id:
            db.execute(
                "INSERT INTO roles (person_id, company_id, title, work_email, work_phone, is_current) VALUES (?,?,?,?,?,?)",
                (person_id, company_id, title, work_email, work_phone, is_current),
            )
    db.commit()


def _get_person_snapshot(db, person_id: int) -> dict:
    """Get current person data for diff comparison."""
    p = db.execute("SELECT * FROM persons WHERE id = ?", (person_id,)).fetchone()
    if not p:
        return {}
    return {k: p[k] for k in p.keys()}


def _compute_diff(old: dict, new: dict) -> str:
    """Compute JSON diff between old and new values."""
    changes = {}
    for key in set(list(old.keys()) + list(new.keys())):
        old_val = old.get(key)
        new_val = new.get(key)
        if old_val != new_val and key not in ("updated_at", "updated_by"):
            changes[key] = {"old": old_val, "new": new_val}
    return json.dumps(changes, ensure_ascii=False) if changes else ""


@router.get("")
@login_required
async def list_persons(request: Request):
    tag_filter = request.query_params.get("tag", "")
    company_filter = request.query_params.get("company", "")
    tier_filter = request.query_params.get("tier", "")
    fam_min = request.query_params.get("fam_min", "")
    page = int(request.query_params.get("page", 1))
    offset = (page - 1) * PER_PAGE
    db = get_db()

    conditions = []
    params = []

    if tag_filter:
        conditions.append("p.id IN (SELECT pt.person_id FROM person_tags pt JOIN tags t ON t.id = pt.tag_id WHERE t.name = ?)")
        params.append(tag_filter)

    if company_filter:
        conditions.append("p.id IN (SELECT r.person_id FROM roles r JOIN companies c ON c.id = r.company_id WHERE c.name = ?)")
        params.append(company_filter)

    if tier_filter:
        conditions.append("p.gift_tier = ?")
        params.append(tier_filter)

    if fam_min:
        conditions.append("p.familiarity >= ?")
        params.append(int(fam_min))

    where = "WHERE " + " AND ".join(conditions) if conditions else ""

    total = db.execute(f"SELECT COUNT(DISTINCT p.id) FROM persons p {where}", params).fetchone()[0]
    total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)

    persons = db.execute(f"""
        SELECT DISTINCT p.* FROM persons p {where}
        ORDER BY p.updated_at DESC LIMIT ? OFFSET ?
    """, params + [PER_PAGE, offset]).fetchall()

    person_list = []
    for p in persons:
        pid = p["id"]
        tags = db.execute("SELECT t.name FROM tags t JOIN person_tags pt ON pt.tag_id = t.id WHERE pt.person_id = ?", (pid,)).fetchall()
        roles = db.execute("SELECT r.*, c.name as company_name FROM roles r LEFT JOIN companies c ON c.id = r.company_id WHERE r.person_id = ?", (pid,)).fetchall()
        person_list.append({"person": dict(p), "tags": [t["name"] for t in tags], "roles": [dict(r) for r in roles]})

    sort_mode = request.query_params.get("sort", "natural")

    if sort_mode == "natural":
        def natural_key(item):
            p = item["person"]
            primary_tag = item["tags"][0] if item["tags"] else ""
            main_role = item["roles"][0] if item["roles"] else {}
            company_name = main_role.get("company_name", "") or ""
            title = main_role.get("title", "") or ""
            fam = p.get("familiarity") or 0
            return (
                industry_priority(primary_tag),
                0 if primary_tag else 1,
                primary_tag,
                company_name,
                title_priority(title),
                -fam,
            )
        person_list.sort(key=natural_key)
    elif sort_mode == "name":
        person_list.sort(key=lambda x: x["person"].get("first_name", ""))

    all_tags = db.execute("SELECT DISTINCT t.name FROM tags t JOIN person_tags pt ON pt.tag_id = t.id ORDER BY t.name").fetchall()
    all_companies = db.execute("SELECT DISTINCT c.name FROM companies c ORDER BY c.name").fetchall()
    db.close()
    return _render(request, "persons/list.html",
        person_list=person_list, all_tags=[t["name"] for t in all_tags],
        all_companies=[c["name"] for c in all_companies],
        current_tag=tag_filter, current_company=company_filter,
        current_tier=tier_filter, current_fam_min=fam_min,
        current_sort=sort_mode,
        page=page, total_pages=total_pages, total=total)


@router.get("/new")
@login_required
async def new_person(request: Request):
    if not can_edit(request.state.user):
        return RedirectResponse("/persons", status_code=302)
    db = get_db()
    companies = db.execute("SELECT id, name FROM companies ORDER BY name").fetchall()
    custom_fields = db.execute("SELECT id, field_name, field_label, field_type, '' as value FROM custom_field_defs ORDER BY sort_order, id").fetchall()
    db.close()
    return _render(request, "persons/form.html",
        person=None, roles=[], tags=[], companies=[dict(c) for c in companies], custom_fields=custom_fields)


def _find_duplicates(db, name: str, email: str, phone: str) -> list:
    """Check for existing contacts that may be duplicates."""
    dupes = []
    seen_ids = set()
    if name:
        rows = db.execute(
            "SELECT id, first_name, email, phone FROM persons WHERE first_name = ?", (name,)
        ).fetchall()
        for r in rows:
            if r["id"] not in seen_ids:
                dupes.append({"person": dict(r), "match": "姓名"})
                seen_ids.add(r["id"])
    if email:
        rows = db.execute(
            "SELECT id, first_name, email, phone FROM persons WHERE email = ?", (email,)
        ).fetchall()
        for r in rows:
            if r["id"] not in seen_ids:
                dupes.append({"person": dict(r), "match": "Email"})
                seen_ids.add(r["id"])
    if phone:
        rows = db.execute(
            "SELECT id, first_name, email, phone FROM persons WHERE phone = ?", (phone,)
        ).fetchall()
        for r in rows:
            if r["id"] not in seen_ids:
                dupes.append({"person": dict(r), "match": "電話"})
                seen_ids.add(r["id"])
    return dupes


@router.post("/new")
@login_required
async def create_person(request: Request):
    if not can_edit(request.state.user):
        return RedirectResponse("/persons", status_code=302)
    form = await request.form()
    db = get_db()

    name = form.get("first_name", "").strip()
    email = form.get("email", "").strip()
    phone = form.get("phone", "").strip()
    force = form.get("force_create", "")

    if not force:
        dupes = _find_duplicates(db, name, email, phone)
        if dupes:
            db.close()
            return _render(request, "persons/duplicate_warning.html",
                duplicates=dupes, form_data=dict(form))

    db.execute(
        "INSERT INTO persons (first_name, last_name, email, phone, notes, gift_tier, birthday, preferences, gift_notes, industry, familiarity, created_by, updated_by) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (form.get("first_name", ""), form.get("last_name", ""), form.get("email", ""),
         form.get("phone", ""), form.get("notes", ""),
         form.get("gift_tier", ""), form.get("birthday", ""),
         form.get("preferences", ""), form.get("gift_notes", ""),
         form.get("industry", ""),
         int(form.get("familiarity")) if form.get("familiarity") else None,
         request.state.user["id"], request.state.user["id"]),
    )
    db.commit()
    person_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]

    _save_tags(db, person_id, form.get("tags", ""))
    _save_roles(db, person_id, form)

    from app.routes.custom_field_routes import save_custom_fields_for_person
    save_custom_fields_for_person(db, person_id, form)

    db.execute(
        "INSERT INTO edit_log (user_id, entity_type, entity_id, action, changes) VALUES (?,?,?,?,?)",
        (request.state.user["id"], "person", person_id, "create", f"建立聯絡人: {name}"),
    )
    db.commit()
    notify_other_editors(db, request.state.user["id"],
        f"{request.state.user['username']} 新增了聯絡人「{name}」", f"/persons/{person_id}")
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


@router.get("/duplicates")
@login_required
async def duplicates_scanner(request: Request):
    db = get_db()
    # Group by normalized name
    persons = db.execute("""
        SELECT p.*,
            (SELECT GROUP_CONCAT(c.name, ', ') FROM roles r JOIN companies c ON c.id = r.company_id WHERE r.person_id = p.id AND r.is_current = 1) as company_names,
            (SELECT r.title FROM roles r WHERE r.person_id = p.id AND r.is_current = 1 LIMIT 1) as current_title
        FROM persons p ORDER BY p.first_name
    """).fetchall()
    db.close()

    from collections import defaultdict
    groups = defaultdict(list)
    for p in persons:
        key = (p["first_name"] or "").strip().lower() + "|" + (p["last_name"] or "").strip().lower()
        groups[key].append(dict(p))

    duplicate_groups = [(k, v) for k, v in groups.items() if len(v) >= 2]
    total_records = sum(len(v) for _, v in duplicate_groups)

    return _render(request, "persons/duplicates.html",
        duplicate_groups=duplicate_groups, group_count=len(duplicate_groups),
        total_records=total_records)


@router.get("/{person_id}")
@login_required
async def view_person(request: Request, person_id: int):
    db = get_db()
    person = db.execute("SELECT * FROM persons WHERE id = ?", (person_id,)).fetchone()
    if not person:
        db.close()
        return RedirectResponse("/persons", status_code=302)

    roles = db.execute("""
        SELECT r.*, c.name as company_name FROM roles r
        LEFT JOIN companies c ON c.id = r.company_id WHERE r.person_id = ? ORDER BY r.is_current DESC
    """, (person_id,)).fetchall()

    tags = db.execute("SELECT t.name FROM tags t JOIN person_tags pt ON pt.tag_id = t.id WHERE pt.person_id = ?", (person_id,)).fetchall()

    relationships = db.execute("""
        SELECT rel.*,
            pa.first_name as person_a_name, pa.id as person_a_id,
            pb.first_name as person_b_name, pb.id as person_b_id
        FROM relationships rel
        JOIN persons pa ON pa.id = rel.person_a_id
        JOIN persons pb ON pb.id = rel.person_b_id
        WHERE rel.person_a_id = ? OR rel.person_b_id = ?
    """, (person_id, person_id)).fetchall()

    interactions = db.execute("""
        SELECT i.*, u.username FROM interactions i
        LEFT JOIN users u ON u.id = i.user_id
        WHERE i.person_id = ? ORDER BY i.interaction_date DESC, i.created_at DESC
    """, (person_id,)).fetchall()

    gift_records = db.execute("""
        SELECT g.*, u.username FROM gift_records g
        LEFT JOIN users u ON u.id = g.user_id
        WHERE g.person_id = ? ORDER BY g.gift_date DESC
    """, (person_id,)).fetchall()

    gift_total = db.execute(
        "SELECT COALESCE(SUM(amount), 0) FROM gift_records WHERE person_id = ?", (person_id,)
    ).fetchone()[0]

    logs = db.execute("""
        SELECT el.*, u.username FROM edit_log el
        LEFT JOIN users u ON u.id = el.user_id
        WHERE el.entity_type = 'person' AND el.entity_id = ?
        ORDER BY el.created_at DESC LIMIT 10
    """, (person_id,)).fetchall()

    all_persons = db.execute("SELECT id, first_name, last_name FROM persons WHERE id != ? ORDER BY first_name", (person_id,)).fetchall()

    # Custom fields
    from app.routes.custom_field_routes import get_custom_fields_for_person
    custom_fields = get_custom_fields_for_person(db, person_id)

    # Attachments
    attachments = db.execute(
        "SELECT * FROM attachments WHERE person_id = ? ORDER BY created_at DESC", (person_id,)
    ).fetchall()

    db.close()
    return _render(request, "persons/detail.html",
        person=person, roles=roles, tags=[t["name"] for t in tags],
        relationships=relationships, interactions=interactions,
        gift_records=gift_records, gift_total=gift_total,
        logs=logs, all_persons=all_persons, custom_fields=custom_fields,
        attachments=attachments)


@router.get("/{person_id}/edit")
@login_required
async def edit_person(request: Request, person_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse(f"/persons/{person_id}", status_code=302)
    db = get_db()
    person = db.execute("SELECT * FROM persons WHERE id = ?", (person_id,)).fetchone()
    if not person:
        db.close()
        return RedirectResponse("/persons", status_code=302)
    roles = db.execute("SELECT r.*, c.name as company_name FROM roles r LEFT JOIN companies c ON c.id = r.company_id WHERE r.person_id = ?", (person_id,)).fetchall()
    tags = db.execute("SELECT t.name FROM tags t JOIN person_tags pt ON pt.tag_id = t.id WHERE pt.person_id = ?", (person_id,)).fetchall()
    companies = db.execute("SELECT id, name FROM companies ORDER BY name").fetchall()
    from app.routes.custom_field_routes import get_custom_fields_for_person
    custom_fields = get_custom_fields_for_person(db, person_id)
    db.close()
    return _render(request, "persons/form.html",
        person=dict(person), roles=[dict(r) for r in roles],
        tags=[t["name"] for t in tags], companies=[dict(c) for c in companies], custom_fields=custom_fields)


@router.post("/{person_id}/edit")
@login_required
async def update_person(request: Request, person_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse(f"/persons/{person_id}", status_code=302)
    form = await request.form()
    db = get_db()

    old_data = _get_person_snapshot(db, person_id)

    db.execute(
        """UPDATE persons SET first_name=?, last_name=?, email=?, phone=?, notes=?,
           gift_tier=?, birthday=?, preferences=?, gift_notes=?, industry=?, familiarity=?,
           updated_by=?, updated_at=CURRENT_TIMESTAMP WHERE id=?""",
        (form.get("first_name", ""), form.get("last_name", ""), form.get("email", ""),
         form.get("phone", ""), form.get("notes", ""),
         form.get("gift_tier", ""), form.get("birthday", ""),
         form.get("preferences", ""), form.get("gift_notes", ""),
         form.get("industry", ""),
         int(form.get("familiarity")) if form.get("familiarity") else None,
         request.state.user["id"], person_id),
    )
    db.commit()

    new_data = _get_person_snapshot(db, person_id)
    diff = _compute_diff(old_data, new_data)

    _save_tags(db, person_id, form.get("tags", ""))
    _save_roles(db, person_id, form)

    from app.routes.custom_field_routes import save_custom_fields_for_person
    save_custom_fields_for_person(db, person_id, form)

    name = form.get('first_name', '')
    db.execute(
        "INSERT INTO edit_log (user_id, entity_type, entity_id, action, changes, diff_json) VALUES (?,?,?,?,?,?)",
        (request.state.user["id"], "person", person_id, "update", f"更新聯絡人: {name}", diff),
    )
    db.commit()
    notify_other_editors(db, request.state.user["id"],
        f"{request.state.user['username']} 更新了聯絡人「{name}」", f"/persons/{person_id}")
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


@router.post("/{person_id}/delete")
@login_required
async def delete_person(request: Request, person_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse(f"/persons/{person_id}", status_code=302)
    db = get_db()
    person = db.execute("SELECT first_name, last_name FROM persons WHERE id = ?", (person_id,)).fetchone()
    name = person['first_name'] if person else ""
    db.execute("DELETE FROM persons WHERE id = ?", (person_id,))
    db.execute(
        "INSERT INTO edit_log (user_id, entity_type, entity_id, action, changes) VALUES (?,?,?,?,?)",
        (request.state.user["id"], "person", person_id, "delete", f"刪除聯絡人: {name}"),
    )
    db.commit()
    notify_other_editors(db, request.state.user["id"],
        f"{request.state.user['username']} 刪除了聯絡人「{name}」", "/persons")
    db.close()
    return RedirectResponse("/persons", status_code=302)


@router.post("/{person_id}/photo")
@login_required
async def upload_photo(request: Request, person_id: int, photo: UploadFile = File(...)):
    if not can_edit(request.state.user):
        return RedirectResponse(f"/persons/{person_id}", status_code=302)
    content = await photo.read()
    ext = os.path.splitext(photo.filename or "photo.jpg")[1] or ".jpg"
    filename = f"person_{person_id}{ext}"
    save_path = PHOTO_DIR / filename
    with open(save_path, "wb") as f:
        f.write(content)
    db = get_db()
    db.execute("UPDATE persons SET photo_url = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
               (f"/static/photos/{filename}", person_id))
    db.commit()
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


@router.post("/{person_id}/relationship")
@login_required
async def add_relationship(request: Request, person_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse(f"/persons/{person_id}", status_code=302)
    form = await request.form()
    other_id = form.get("other_person_id")
    label = form.get("label", "")
    notes = form.get("rel_notes", "")

    if other_id and label:
        db = get_db()
        db.execute("INSERT INTO relationships (person_a_id, person_b_id, label, notes) VALUES (?,?,?,?)",
                   (person_id, int(other_id), label, notes))
        db.commit()
        db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


@router.post("/relationship/{rel_id}/delete")
@login_required
async def delete_relationship(request: Request, rel_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse("/persons", status_code=302)
    form = await request.form()
    person_id = form.get("person_id", "")
    db = get_db()
    db.execute("DELETE FROM relationships WHERE id = ?", (rel_id,))
    db.commit()
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


# --- Interaction routes ---
@router.post("/{person_id}/interaction")
@login_required
async def add_interaction(request: Request, person_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse(f"/persons/{person_id}", status_code=302)
    form = await request.form()
    db = get_db()
    db.execute(
        "INSERT INTO interactions (person_id, user_id, type, title, content, interaction_date) VALUES (?,?,?,?,?,?)",
        (person_id, request.state.user["id"], form.get("type", "note"),
         form.get("title", ""), form.get("content", ""), form.get("interaction_date", "")),
    )
    db.commit()
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


@router.post("/interaction/{interaction_id}/delete")
@login_required
async def delete_interaction(request: Request, interaction_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse("/persons", status_code=302)
    form = await request.form()
    person_id = form.get("person_id", "")
    db = get_db()
    db.execute("DELETE FROM interactions WHERE id = ?", (interaction_id,))
    db.commit()
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


# --- Gift record routes ---
@router.post("/{person_id}/gift")
@login_required
async def add_gift_record(request: Request, person_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse(f"/persons/{person_id}", status_code=302)
    form = await request.form()
    db = get_db()
    amount = 0
    try:
        amount = float(form.get("amount", 0))
    except (ValueError, TypeError):
        pass
    db.execute(
        "INSERT INTO gift_records (person_id, user_id, gift_date, occasion, item, amount, notes) VALUES (?,?,?,?,?,?,?)",
        (person_id, request.state.user["id"], form.get("gift_date", ""),
         form.get("occasion", ""), form.get("item", ""),
         amount, form.get("gift_record_notes", "")),
    )
    db.commit()
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


@router.post("/gift/{gift_id}/delete")
@login_required
async def delete_gift_record(request: Request, gift_id: int):
    if not can_edit(request.state.user):
        return RedirectResponse("/persons", status_code=302)
    form = await request.form()
    person_id = form.get("person_id", "")
    db = get_db()
    db.execute("DELETE FROM gift_records WHERE id = ?", (gift_id,))
    db.commit()
    db.close()
    return RedirectResponse(f"/persons/{person_id}", status_code=302)


# --- Batch operations ---
@router.post("/batch")
@login_required
async def batch_action(request: Request):
    if not can_edit(request.state.user):
        return RedirectResponse("/persons", status_code=302)
    form = await request.form()
    action = form.get("batch_action", "")
    person_ids = form.getlist("selected[]")

    if not person_ids:
        return RedirectResponse("/persons", status_code=302)

    db = get_db()

    if action == "export_selected":
        import io
        import openpyxl
        from fastapi.responses import StreamingResponse

        ids = [int(pid) for pid in person_ids]
        placeholders = ",".join(["?"] * len(ids))
        persons = db.execute(f"SELECT * FROM persons WHERE id IN ({placeholders}) ORDER BY first_name", ids).fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "貴賓選取名單"
        headers = ['姓名', '英文名', '公司', '職稱', '電話', 'Email', '地址', '產業', '熟識程度', '備註']
        ws.append(headers)

        for p in persons:
            pid_val = p["id"]
            role = db.execute("""
                SELECT r.title, c.name as company_name, c.address
                FROM roles r LEFT JOIN companies c ON c.id = r.company_id
                WHERE r.person_id = ? AND r.is_current = 1 LIMIT 1
            """, (pid_val,)).fetchone()
            ws.append([
                p["first_name"] or "",
                p["last_name"] or "",
                role["company_name"] if role else "",
                role["title"] if role else "",
                p["phone"] or "",
                p["email"] or "",
                role["address"] if role and role["address"] else "",
                p.get("industry", "") or "",
                f"{'★' * p['familiarity']}" if p.get("familiarity") else "",
                p["notes"] or "",
            ])

        db.close()
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"貴賓選取名單（{len(persons)}位）.xlsx"

        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
        )

    if action == "delete":
        for pid in person_ids:
            db.execute("DELETE FROM persons WHERE id = ?", (int(pid),))
        db.commit()
        request.state.flash = [("success", f"已刪除 {len(person_ids)} 位聯絡人")]

    elif action == "add_tag":
        tag_name = form.get("batch_tag", "").strip()
        if tag_name:
            tag_id = _get_or_create_tag(db, tag_name)
            for pid in person_ids:
                db.execute("INSERT OR IGNORE INTO person_tags (person_id, tag_id) VALUES (?,?)", (int(pid), tag_id))
            db.commit()
            request.state.flash = [("success", f"已為 {len(person_ids)} 位聯絡人加上標籤「{tag_name}」")]

    elif action == "remove_tag":
        tag_name = form.get("batch_tag", "").strip().lower()
        if tag_name:
            tag = db.execute("SELECT id FROM tags WHERE name = ?", (tag_name,)).fetchone()
            if tag:
                for pid in person_ids:
                    db.execute("DELETE FROM person_tags WHERE person_id = ? AND tag_id = ?", (int(pid), tag["id"]))
                db.commit()
            request.state.flash = [("success", f"已為 {len(person_ids)} 位聯絡人移除標籤「{tag_name}」")]

    db.close()
    return RedirectResponse("/persons", status_code=302)
